import csv
import io
from pathlib import Path

from pypdf import PdfReader
from docx import Document
import openpyxl
from pptx import Presentation


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".pdf",
    ".docx",
    ".doc",
    ".csv",
    ".xlsx",
    ".xls",
    ".pptx",
    ".ppt",
}


def load_text_content(stream_or_path) -> str:
    if isinstance(stream_or_path, (str, Path)):
        p = Path(stream_or_path)
        try:
            return p.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            return p.read_text(encoding="latin-1")
    elif isinstance(stream_or_path, bytes):
        try:
            return stream_or_path.decode("utf-8")
        except UnicodeDecodeError:
            return stream_or_path.decode("latin-1")
    elif hasattr(stream_or_path, "read"):
        raw = stream_or_path.read()
        if isinstance(raw, str):
            return raw
        try:
            return raw.decode("utf-8")
        except UnicodeDecodeError:
            return raw.decode("latin-1")
    return ""


def load_pdf_file(source) -> str:
    if isinstance(source, bytes):
        reader = PdfReader(io.BytesIO(source))
    elif hasattr(source, "read"):
        reader = PdfReader(source)
    else:
        reader = PdfReader(str(source))

    pages = []
    for idx, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"--- Page {idx} ---\n{text.strip()}")

    return "\n\n".join(pages)


def load_docx_file(source) -> str:
    if isinstance(source, bytes):
        source = io.BytesIO(source)
    doc = Document(source)

    parts = []
    # Extract paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            parts.append(txt)

    # Extract tables
    for table_idx, table in enumerate(doc.tables, start=1):
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(row_cells):
                table_rows.append(" | ".join(row_cells))
        if table_rows:
            parts.append(f"[Table {table_idx}]\n" + "\n".join(table_rows))

    return "\n\n".join(parts)


def load_csv_file(source) -> str:
    raw_text = load_text_content(source)
    reader = csv.reader(io.StringIO(raw_text))
    rows = list(reader)
    if not rows:
        return ""

    headers = rows[0]
    lines = [f"Columns: {', '.join(headers)}"]
    for idx, row in enumerate(rows[1:], start=1):
        if any(cell.strip() for cell in row):
            record = [f"{headers[i] if i < len(headers) else f'col_{i}'}: {val.strip()}"
                      for i, val in enumerate(row) if val.strip()]
            lines.append(f"Row {idx} -> {'; '.join(record)}")

    return "\n".join(lines)


def load_excel_file(source) -> str:
    if isinstance(source, bytes):
        source = io.BytesIO(source)
    wb = openpyxl.load_workbook(source, data_only=True)
    sheets_content = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        sheet_lines = [f"=== Sheet: {sheet_name} ==="]
        header = None
        for r_idx, row in enumerate(rows):
            filtered = [str(cell).strip() if cell is not None else "" for cell in row]
            if not any(filtered):
                continue
            if header is None:
                header = filtered
                sheet_lines.append(f"Columns: {', '.join(header)}")
            else:
                row_str = [f"{header[i] if i < len(header) else f'col_{i}'}: {val}"
                           for i, val in enumerate(filtered) if val]
                sheet_lines.append(f"Row {r_idx} -> {'; '.join(row_str)}")

        sheets_content.append("\n".join(sheet_lines))

    return "\n\n".join(sheets_content)


def load_pptx_file(source) -> str:
    if isinstance(source, bytes):
        source = io.BytesIO(source)
    prs = Presentation(source)
    slides_text = []

    for idx, slide in enumerate(prs.slides, start=1):
        slide_parts = [f"--- Slide {idx} ---"]
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_parts.append(shape.text.strip())

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_parts.append(f"Speaker Notes: {notes}")

        if len(slide_parts) > 1:
            slides_text.append("\n".join(slide_parts))

    return "\n\n".join(slides_text)


def load_document(file_path: Path) -> str:
    suffix = file_path.suffix.lower()

    if suffix in {".txt", ".md"}:
        return load_text_content(file_path)

    if suffix == ".pdf":
        return load_pdf_file(file_path)

    if suffix in {".docx", ".doc"}:
        return load_docx_file(file_path)

    if suffix == ".csv":
        return load_csv_file(file_path)

    if suffix in {".xlsx", ".xls"}:
        return load_excel_file(file_path)

    if suffix in {".pptx", ".ppt"}:
        return load_pptx_file(file_path)

    raise ValueError(
        f"Unsupported file type: {suffix}. "
        f"Supported types: {SUPPORTED_EXTENSIONS}"
    )


def load_document_from_bytes(file_bytes: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()

    if suffix in {".txt", ".md"}:
        return load_text_content(file_bytes)

    if suffix == ".pdf":
        return load_pdf_file(file_bytes)

    if suffix in {".docx", ".doc"}:
        return load_docx_file(file_bytes)

    if suffix == ".csv":
        return load_csv_file(file_bytes)

    if suffix in {".xlsx", ".xls"}:
        return load_excel_file(file_bytes)

    if suffix in {".pptx", ".ppt"}:
        return load_pptx_file(file_bytes)

    raise ValueError(
        f"Unsupported file type: {suffix}. "
        f"Supported types: {SUPPORTED_EXTENSIONS}"
    )


def load_documents(directory: str) -> list[dict]:
    directory_path = Path(directory)

    if not directory_path.exists():
        raise FileNotFoundError(
            f"Document directory does not exist: {directory}"
        )

    documents = []

    for file_path in sorted(directory_path.iterdir()):
        if not file_path.is_file():
            continue

        if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        try:
            text = load_document(file_path)
        except Exception as e:
            print(f"Warning: Failed to load {file_path.name}: {e}")
            continue

        if not text.strip():
            continue

        documents.append(
            {
                "document_id": file_path.stem,
                "file_name": file_path.name,
                "file_path": str(file_path),
                "text": text.strip(),
            }
        )

    return documents
