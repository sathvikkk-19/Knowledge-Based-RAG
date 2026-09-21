import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from docx import Document
from pptx import Presentation

from app.ingestion.loader import (
    load_document_from_bytes,
    load_csv_file,
    load_docx_file,
    load_excel_file,
    load_pptx_file,
    load_text_content,
    SUPPORTED_EXTENSIONS,
)
from app.ingestion.embedder import Embedder
from app.ingestion.pipeline import sanitize_document_id


def test_supported_extensions():
    print("Testing supported extensions...")
    expected = {".pdf", ".docx", ".doc", ".csv", ".xlsx", ".xls", ".pptx", ".ppt", ".txt", ".md"}
    assert expected.issubset(SUPPORTED_EXTENSIONS), f"Missing extensions: {expected - SUPPORTED_EXTENSIONS}"
    print("  -> Extensions check passed!")


def test_csv_loader():
    print("Testing CSV loader...")
    csv_bytes = b"Product,Category,Price\nCloudOS,Infrastructure,199\nDataMesh,Database,299"
    text = load_csv_file(csv_bytes)
    assert "Columns: Product, Category, Price" in text
    assert "CloudOS" in text and "DataMesh" in text
    print("  -> CSV loader passed!")


def test_docx_loader():
    print("Testing DOCX loader...")
    doc = Document()
    doc.add_heading("Cloud Architecture", level=1)
    doc.add_paragraph("Acme Platform uses PostgreSQL for storage.")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Neo4j"
    table.cell(1, 1).text = "Healthy"

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    text = load_docx_file(buf.getvalue())
    assert "Cloud Architecture" in text
    assert "PostgreSQL" in text
    assert "Neo4j" in text
    print("  -> DOCX loader passed!")


def test_excel_loader():
    print("Testing Excel loader...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Services"
    ws.append(["Service", "Port", "Protocol"])
    ws.append(["Neo4j", 7687, "Bolt"])
    ws.append(["Postgres", 5432, "TCP"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    text = load_excel_file(buf.getvalue())
    assert "Services" in text
    assert "7687" in text
    assert "Postgres" in text
    print("  -> Excel loader passed!")


def test_pptx_loader():
    print("Testing PowerPoint loader...")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    title.text = "System Overview"
    subtitle = slide.placeholders[1]
    subtitle.text = "Knowledge Graph RAG Architecture"

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)

    text = load_pptx_file(buf.getvalue())
    assert "System Overview" in text
    assert "Knowledge Graph RAG Architecture" in text
    print("  -> PowerPoint loader passed!")


def test_embedder_dimension():
    print("Testing Embedder output dimensions...")
    embedder = Embedder()
    vec = embedder.embed("Knowledge Graph RAG platform")
    assert len(vec) == 384, f"Expected 384 dimensions, got {len(vec)}"
    print(f"  -> Embedder passed (384 dimensions using {embedder.backend})!")


def test_sanitization():
    print("Testing document ID sanitization...")
    assert sanitize_document_id("My Report (2026) Final.docx") == "my_report_2026_final"
    assert sanitize_document_id("data-sheet_v2.csv") == "data_sheet_v2"
    print("  -> Sanitization passed!")


if __name__ == "__main__":
    test_supported_extensions()
    test_csv_loader()
    test_docx_loader()
    test_excel_loader()
    test_pptx_loader()
    test_embedder_dimension()
    test_sanitization()
    print("\nALL MULTI-FORMAT UPLOAD TESTS PASSED SUCCESSFULLY!")
